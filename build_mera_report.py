#!/usr/bin/env python3
"""
Reconstruit le rapport de stage MERA depuis le modèle ENSPD (DOMCHE).
- Conserve la page de garde (avec nom/axe modifiés) et le Chapitre 1 à l'identique.
- Met à jour l'introduction, le chapitre 2, le chapitre 3 et la conclusion.
- Remplace les figures (ESP32, capteurs, RFID, captures de l'application MERA).
"""
import os
import shutil
import docx
from docx.oxml.ns import qn

SRC = '/home/darelle/Documents/MERA_APP/Rapport_de_stage_ENSPD (DOMCHE).docx'
DST = '/home/darelle/Documents/MERA_APP/Rapport_de_stage_ENSPD (MERA).docx'


def normalize_quotes(text):
    """Remplace les apostrophes droites par des apostrophes typographiques."""
    return text.replace("'", "\u2019")


def replace_text_only(para, text):
    """Remplace uniquement le texte, en préservant les images/dessins du paragraphe."""
    text = normalize_quotes(text)
    text_runs = []
    for r in para.runs:
        el = r._element
        if (el.findall(qn('w:drawing')) or el.findall(qn('w:pict'))
                or el.findall(qn('w:object'))):
            continue
        text_runs.append(r)
    if not text_runs:
        return False
    text_runs[0].text = text
    for r in text_runs[1:]:
        r.text = ''
    return True


# ============================================================
# CONTENU
# ============================================================

TDM = {
    105: 'DÉDICACES\ti',
    106: 'REMERCIEMENTS\tii',
    107: 'TABLE DE MATIERES\tiii',
    108: 'LISTES DES FIGURES\tiv',
    109: 'INTRODUCTION\t1',
    110: 'Chapitre 1 : PRÉSENTATION DE L\'ENTREPRISE\t2',
    111: '1. Plan de localisation\t4',
    112: '2. Organigramme de l\'entreprise\t5',
    113: 'Chapitre 2 : TRAVAUX EFFECTUES AU COURS DU STAGE\t6',
    114: '1. Les travaux effectués\t6',
    115: '2. Les outils mis à ma disposition\t9',
    116: 'a. Outils matériels\t9',
    117: 'b. Outils logiciels\t11',
    118: '3. Les Missions du poste occupé\t13',
    119: '4. Taches périphériques\t14',
    120: 'Chapitre 3 : ACQUIS DU STAGE\t15',
    121: '1. Compétences acquises\t15',
    122: '2. Difficultés rencontrées et solutions apportées\t17',
    123: '3. La vie en société\t18',
    124: 'CONCLUSION\t19',
}

FIGURES_LIST = {
    132: 'Figure 1: Carte de développement ESP32\t9',
    133: 'Figure 2: Capteur cardiaque MAX30102\t9',
    134: 'Figure 3: Capteur de température MLX90614\t10',
    135: 'Figure 4: Module RFID MFRC522\t10',
    136: 'Figure 5: Balance numérique intelligente (Famkit)\t11',
    137: 'Figure 7: Logo TypeScript\t11',
    138: 'Figure 8: Logo Node.js\t11',
    139: 'Figure 9: Logo PostgreSQL\t11',
    140: 'Figure 10: Logo Prisma\t12',
    141: 'Figure 11: Tableau de bord de l\'application MERA\t15',
}

INTRO = {
    157: 'Le stage constitue une étape fondamentale dans la formation des ingénieurs, car il permet aux étudiants d\'allier les connaissances théoriques acquises à l\'école aux réalités pratiques du monde professionnel. C\'est également une occasion privilégiée de découvrir le fonctionnement d\'une structure, de comprendre ses missions et son organisation, tout en développant des compétences techniques et humaines indispensables à l\'exercice du métier d\'ingénieur.',
    159: 'C\'est dans ce cadre que s\'inscrit le présent rapport, élaboré à la suite de mon stage effectué à l\'École Nationale Supérieure Polytechnique de Douala (ENSPD). Cet établissement public d\'enseignement supérieur, rattaché à l\'Université de Douala, m\'a offert un cadre stimulant pour travailler sur le projet MERA — Medical Embedded Robotic Assistant — un système médical embarqué intelligent destiné au dépistage et au diagnostic dans les centres de santé, conçu autour du microcontrôleur ESP32 et doté d\'une application web complète de suivi ophtalmologique.',
    161: 'Ce rapport se propose de présenter le déroulement du stage, les activités techniques auxquelles j\'ai pris part — notamment la conception du système embarqué MERA, le développement de son application web full-stack et la réalisation du circuit imprimé (PCB) — ainsi que les enseignements et compétences acquis au sein de cette institution.',
}

CH2 = {
    # 1. Les travaux effectués
    205: 'Au cours de mon stage, j\'ai eu l\'opportunité de travailler sur le projet MERA (Medical Embedded Robotic Assistant), un système médical embarqué intelligent destiné au dépistage et au diagnostic dans les centres de santé. Ce projet, qui combine électronique embarquée, intelligence artificielle et développement web, m\'a permis de mettre en pratique mes connaissances et d\'acquérir de nouvelles compétences techniques grâce à une approche essentiellement pratique et à l\'accompagnement de mes encadreurs.',
    206: 'Les travaux réalisés au cours de ce stage peuvent être regroupés en trois grands axes :',
    207: 'Conception et développement du système embarqué MERA : intégration de multiples capteurs (cardiaque, thermométrique, RFID) et d\'une balance numérique intelligente connectée en Bluetooth Low Energy autour du microcontrôleur ESP32, avec gestion temps réel via FreeRTOS, affichage sur écran TFT et communication WiFi.',
    208: 'Développement de l\'application web full-stack MERA : création d\'une plateforme de diagnostic ophtalmologique combinant un frontend React, un backend Node.js/Express, une base de données Prisma/PostgreSQL et des modules d\'intelligence artificielle pour l\'analyse des maladies oculaires.',
    209: 'Conception du circuit imprimé (PCB) MERA : réalisation du schéma et du routage de la carte électronique sous KiCad, intégrant le module ESP32, les capteurs biométriques, l\'écran TFT et les circuits d\'alimentation.',
    210: 'Ces différents travaux ont constitué une phase d\'apprentissage et de montée en compétence, notamment dans l\'intégration de capteurs biométriques, le développement d\'applications web modernes et la conception de circuits imprimés.',
    211: 'Intégration d\'une balance numérique intelligente connectée en Bluetooth Low Energy (BLE) : le poids du patient est reçu directement par l\'ESP32 à partir des annonces diffusées par la balance (modèle Famkit), validé après stabilisation sur plusieurs lectures consécutives, puis enregistré avec les autres constantes vitales.',
    212: 'Le projet MERA est un système médical embarqué conçu pour assister le personnel soignant dans le dépistage et le diagnostic de pathologies, en particulier oculaires, dans les centres de santé. Le système repose sur un microcontrôleur ESP32 qui orchestre l\'ensemble des capteurs et modules de communication, couplé à une application web permettant le suivi des patients, l\'analyse par intelligence artificielle et la génération de rapports médicaux.',
    213: 'Les travaux réalisés dans le cadre de ce projet incluent :',
    214: 'Conception et mise en œuvre du système embarqué MERA, basé sur le microcontrôleur ESP32, assurant la collecte des données biométriques (rythme cardiaque, SpO2, température corporelle, poids) via des capteurs dédiés, ainsi que l\'identification des patients par badge RFID.',
    215: 'Intégration du capteur cardiaque MAX30102, mesurant le taux d\'oxygène dans le sang (SpO2) et le rythme cardiaque par photopléthysmographie. Le traitement du signal est assuré par un algorithme de type Kalman avec filtrage adaptatif, exécuté sur un cœur du processeur double cœur de l\'ESP32 via FreeRTOS.',
    216: 'Intégration du capteur de température infrarouge MLX90614, permettant une mesure sans contact de la température corporelle avec une précision de ±0.5°C. Les lectures sont validées par un mécanisme de stabilisation sur cinq échantillons.',
    217: 'Intégration du module RFID MFRC522 pour l\'identification des patients et du personnel soignant par badge. La communication avec l\'ESP32 principal est assurée par un second ESP32 esclave via liaison UART, avec protocole de handshake et surveillance de santé.',
    218: 'Développement de l\'interface utilisateur sur écran TFT ILI9488 (480×320 pixels) pilotée par la librairie TFT_eSPI, affichant les résultats de mesure en temps réel et guidant le patient à travers le protocole de dépistage.',
    219: 'Mise en place de la communication WiFi et des échanges avec le serveur backend : envoi des mesures via le protocole HTTP, heartbeat de surveillance, et file d\'attente hors ligne pour la retransmission des données en cas d\'absence de connexion.',
    220: 'Développement de l\'application web MERA, plateforme full-stack combinant un frontend React avec interface de diagnostic, un backend Node.js/Express avec API REST, et une intelligence artificielle multi-fournisseurs (Groq, Gemini, Anthropic) pour l\'analyse automatique des pathologies oculaires.',
    221: 'Conception du circuit imprimé (PCB) MERA sous KiCad, intégrant le module ESP32, les capteurs MAX30102 et MLX90614, l\'écran TFT, un buzzer, un circuit de charge solaire (CN3791) et une batterie lithium-ion pour l\'autonomie du dispositif.',
    222: 'L\'ensemble de ces travaux m\'a permis de développer mes compétences dans le domaine des systèmes embarqués, du traitement du signal biomédical, du développement web full-stack, de la conception de circuits imprimés et de l\'intégration d\'intelligence artificielle dans un contexte médical.',

    # 2. Les outils mis à ma disposition
    224: 'Au cours de mon stage, j\'ai eu accès à un ensemble d\'outils matériels et logiciels indispensables à la réalisation du projet MERA. Ces ressources m\'ont permis d\'acquérir des compétences pratiques et de mener à bien les missions qui m\'ont été confiées.',

    # Outils matériels
    226: 'Microcontrôleur ESP32',
    227: 'Microcontrôleur double cœur 32 bits (ESP32-WROOM) constituant le cœur du système MERA. Il assure la gestion des capteurs via les bus SPI et I2C, le traitement du signal biomédical en temps réel via FreeRTOS, l\'affichage sur écran TFT et la communication WiFi avec le serveur ainsi que la gestion du module RFID associé.',
    231: 'Figure 1: Carte de développement ESP32',
    232: 'Capteur cardiaque MAX30102',
    234: 'Module de mesure cardiaque intégrant deux LED (infrarouge et rouge) et un photodétecteur, permettant la mesure simultanée du taux d\'oxygène dans le sang (SpO2) et du rythme cardiaque par photopléthysmographie. L\'algorithme de traitement du signal utilise un filtre de Kalman et une tolérance adaptative pour assurer la fiabilité des mesures.',
    239: 'Figure 2: Capteur cardiaque MAX30102',
    242: 'Capteur de température MLX90614',
    243: 'Capteur de température infrarouge sans contact, mesurant la température corporelle avec une précision de ±0.5°C sur une plage de -70 à 380°C. Communiquant via le bus I2C, il est utilisé pour le dépistage thermique non invasif des patients.',
    246: 'Figure 3: Capteur de température MLX90614',
    247: 'Module RFID MFRC522',
    248: 'Lecteur de badges RFID 13.56 MHz permettant l\'identification des patients et du personnel soignant. Associé à un second ESP32 esclave communiquant via UART, il assure une identification rapide et fiable sans contact.',
    253: 'Figure 4: Module RFID MFRC522',

    # Outils logiciels
    255: 'Outils logiciels',
    256: 'Arduino IDE : environnement de développement utilisé pour programmer le microcontrôleur ESP32 et l\'ensemble des capteurs. La programmation en C++ avec FreeRTOS permet la gestion multitâche temps réel du système embarqué.',
    258: 'Figure 6: Logo Arduino IDE',
    259: 'FreeRTOS : Système d\'exploitation temps réel utilisé pour organiser le fonctionnement multitâche du firmware. Il permet de séparer la gestion des capteurs (cœur 0) des communications et de l\'affichage (cœur 1), en utilisant des mutex pour la synchronisation des données partagées.',
    260: 'KiCad : logiciel de conception de circuits imprimés utilisé pour la réalisation du PCB MERA. Le schéma intègre le module ESP32, les capteurs biométriques, l\'écran TFT, les circuits d\'alimentation (régulateur 3.3V, chargeur solaire CN3791, chargeur LiPo TP4056) et le buzzer.',
    261: 'Communication WiFi et API REST : l\'ESP32 communique avec le serveur backend via le WiFi en utilisant le protocole HTTP. Les mesures biométriques sont transmises à l\'API REST, avec un heartbeat de surveillance et une file d\'attente hors ligne permettant la retransmission des données en cas de perte de connexion.',
    262: 'Node.js et TypeScript : Node.js constitue l\'environnement d\'exécution du backend de l\'application MERA, développé en TypeScript afin de bénéficier du typage statique et d\'une meilleure maintenabilité du code. L\'API REST gère l\'authentification JWT, le CRUD des entités médicales, l\'upload d\'images et l\'invocation des modèles d\'intelligence artificielle.',
    263: 'Prisma et PostgreSQL : Prisma est utilisé comme ORM pour communiquer avec la base de données PostgreSQL. La base structure les informations relatives aux patients, aux sessions de diagnostic, aux signes vitaux, aux photos oculaires, aux résultats d\'analyse et aux traitements traditionnels.',
    265: 'React et Vite : React est utilisé pour développer l\'interface utilisateur de l\'application MERA, avec Vite comme outil de build rapide. L\'interface comprend des tableaux de bord, des sessions de diagnostic, un système de pharmacopée traditionnelle multilingue et la génération de rapports médicaux PDF.',
    269: 'Intelligence artificielle multi-fournisseurs',
    270: 'Système d\'analyse par IA avec auto-basculement entre Groq (Llama 3.3 70B), Gemini (1.5 Flash) et Anthropic, permettant le diagnostic automatique de huit classes de pathologies oculaires (cataracte, conjonctivite, glaucome, rétinopathie diabétique, ptérygion, trachome, uvéite, état sain) avec un mode de secours pour les démonstrations hors ligne.',

    # 3. Les Missions du poste occupé
    274: 'Comme indiqué dans l\'introduction, le rôle d\'un stagiaire en informatique au sein de l\'École Nationale Supérieure Polytechnique de Douala (ENSPD) est déterminant pour contribuer à la modernisation des infrastructures numériques et au développement de solutions technologiques innovantes. Au cours de mon stage, j\'ai été impliqué dans le projet MERA, couvrant un large éventail de missions techniques :',
    275: 'Conception et développement du firmware ESP32 pour le système embarqué MERA, incluant l\'intégration de capteurs biométriques (MAX30102, MLX90614), la récupération du poids par la balance numérique intelligente en Bluetooth Low Energy (BLE), l\'identification RFID, l\'affichage TFT et la communication WiFi avec le serveur.',
    276: 'Développement de l\'application web full-stack MERA, comprenant le frontend React (interface de diagnostic, tableau de bord, pharmacopée, rapports PDF) et le backend Node.js/Express (API REST, authentification, intelligence artificielle, communication robot).',
    277: 'Conception du circuit imprimé (PCB) MERA sous KiCad, incluant le schéma électrique, le routage de la carte et la préparation des fichiers de fabrication (GERBER).',
    278: 'Intégration de l\'intelligence artificielle pour le diagnostic ophtalmologique, avec mise en place d\'un système multi-fournisseurs (Groq, Gemini, Anthropic) et développement d\'algorithmes de classification des pathologies oculaires.',
    279: 'Développement du système de traitement du signal biomédical, incluant l\'algorithme de détection de battements cardiaques avec filtre de Kalman, le calcul de SpO2 et la stabilisation des mesures de température.',
    280: 'Autoformation et montée en compétences sur de nouvelles technologies (FreeRTOS, conception PCB, intelligence artificielle, développement web React/Node.js) afin de mener à bien les différentes composantes du projet MERA.',

    # 4. Taches périphériques
    282: 'En dehors des missions principales liées au projet MERA, j\'ai également participé à un certain nombre de tâches périphériques qui ont contribué à enrichir mon expérience :',
    283: 'Participation à la réalisation d\'une vidéo de présentation (pitch) du projet MERA, couvrant le problème adressé, la solution proposée, le marché cible et le modèle économique.',
    284: 'Réalisation de documents de présentation du projet (diaporamas PowerPoint) pour les manifestations technologiques et les occasions de communication.',
    285: 'Participation à des réunions d\'équipe afin de mieux cerner la coordination entre les différents départements de l\'ENSPD.',
    286: 'Observation et prise de notes sur le fonctionnement des services de l\'établissement, afin de comprendre les besoins numériques spécifiques à chaque secteur.',
    287: 'Réalisation de petites tâches de maintenance de premier niveau sur les équipements bureautiques (imprimantes, périphériques, câblage réseau simple).',
    288: 'Ces activités m\'ont permis de développer une polyvalence, d\'améliorer ma capacité d\'adaptation et de renforcer mon sens de la collaboration interservices.',
}

CH3 = {
    291: 'Dès les premières semaines de ce stage, nous avons eu l\'opportunité de développer et de consolider un large éventail de compétences, en lien direct avec les différentes composantes du projet MERA — système embarqué, application web et conception PCB. Ces apprentissages se sont construits progressivement, à travers la pratique quotidienne, l\'autoformation, mais aussi les échanges avec l\'équipe technique de l\'établissement. Les compétences acquises au cours de ce stage peuvent être regroupées en deux grands volets : les compétences techniques et les compétences relationnelles et organisationnelles.',

    293: 'Compétences techniques :',
    294: 'Développement de firmware embarqué sur ESP32 avec FreeRTOS, incluant la gestion multitâche temps réel, la synchronisation par mutex et la programmation de capteurs biométriques (MAX30102, MLX90614) ainsi que la lecture du poids via une balance numérique connectée en Bluetooth Low Energy.',
    295: 'Intégration de modules de communication : WiFi, UART (liaison maître-esclave), I2C (capteurs biométriques), SPI (écran TFT, RFID) et Bluetooth Low Energy (balance numérique intelligente).',
    297: 'Figure 11: Tableau de bord de l\'application MERA',
    298: 'Conception et routage de circuits imprimés (PCB) sous KiCad, incluant la réalisation du schéma électrique, le placement des composants et la préparation des fichiers de fabrication GERBER.',
    301: 'Développement d\'applications web full-stack avec React (frontend) et Node.js/Express (backend), incluant l\'authentification JWT, le CRUD d\'entités médicales et l\'intégration d\'intelligence artificielle.',
    302: 'Implémentation d\'algorithmes de traitement du signal biomédical : détection de battements cardiaques par photopléthysmographie avec filtre de Kalman, calcul de SpO2 et stabilisation des mesures de température.',
    303: 'Compétences relationnelles et organisationnelles :',
    304: 'Travail en équipe et collaboration avec les encadreurs et les techniciens pour la mise en place et le suivi des différentes composantes du projet MERA.',
    305: 'Développement de l\'autonomie par l\'autoformation sur de nouvelles technologies (FreeRTOS, KiCad, React, Prisma, intelligence artificielle).',
    306: 'Amélioration des capacités de communication technique, notamment pour la présentation des résultats et la documentation des procédures de développement.',
    307: 'Renforcement de la rigueur, de l\'organisation et de la gestion du temps dans l\'accomplissement de missions multi-composantes (embarqué, web, PCB).',

    309: 'Difficultés rencontrées et solutions apportées',
    310: 'Au cours de ce stage, la réalisation du projet MERA n\'a pas été exempte de difficultés. Chaque composante (système embarqué, application web, PCB) a soulevé des défis techniques qu\'il a fallu analyser et surmonter afin d\'atteindre les objectifs fixés. Ces obstacles se sont révélés formatifs, puisqu\'ils ont permis d\'acquérir de nouvelles compétences et de développer des capacités d\'adaptation.',
    312: 'Difficultés techniques',
    313: 'Système embarqué ESP32 et capteurs : L\'intégration simultanée du capteur cardiaque MAX30102, du capteur de température MLX90614 et de l\'écran TFT sur des bus de communication différents (I2C et SPI) a posé des défis de synchronisation. L\'utilisation de FreeRTOS avec des mutex et des tâches réparties sur chaque cœur a résolu ces problèmes de concurrence.',
    314: 'Identification RFID et balance BLE avec double ESP32 : La communication entre l\'ESP32 maître et l\'ESP32 esclave via UART, avec protocole de handshake et surveillance de santé, ainsi que la réception du poids par les annonces Bluetooth Low Energy de la balance numérique (validation sur plusieurs lectures stables et filtrage des poids non physiologiques), ont nécessité une mise au point minutieuse pour garantir la fiabilité des mesures et la détection des déconnexions.',
    315: 'Traitement du signal cardiaque : Le bruit inhérent aux mesures photopléthysmographiques du MAX30102 a nécessité l\'implémentation d\'un algorithme de filtrage par Kalman avec tolérance adaptative, d\'un seuil de détection de doigt et d\'un mécanisme de stabilisation sur plusieurs battements.',
    316: 'Application web et intelligence artificielle : L\'intégration de plusieurs fournisseurs IA (Groq, Gemini, Anthropic) avec mécanisme d\'auto-basculement a posé des défis de compatibilité des formats de réponse. La mise en place d\'un mode de secours pour les démonstrations hors ligne a permis de surmonter ces contraintes.',
    317: 'Conception PCB sous KiCad : Le routage du circuit imprimé intégrant le module ESP32, les capteurs, l\'écran TFT et les circuits d\'alimentation a nécessité une attention particulière au plan de masse, à la séparation des plans d\'alimentation et au respect des largeurs de pistes minimales.',
    318: 'Gestion de l\'alimentation : L\'intégration du circuit de charge solaire (CN3791), du chargeur LiPo (TP4056) et du régulateur 3.3V (LD1117) a posé des problèmes de stabilité de tension lors des pics de courant consommés par le module WiFi et le buzzer. L\'ajout de condensateurs de découplage a résolu ces problèmes.',
    319: 'Communication WiFi et file d\'attente hors ligne : La mise en place d\'un système de file d\'attente pour la retransmission des mesures en cas d\'absence de connexion WiFi a nécessité la gestion de la mémoire limitée de l\'ESP32 et la conception d\'un protocole de reconnexion automatique.',
    320: 'Multilingue et accessibilité : Le développement de l\'interface utilisateur en quatre langues (français, anglais, pulaar, ewondo) avec système de traduction a représenté un défi d\'organisation du code et de gestion des variations linguistiques pour les termes médicaux.',
    321: 'Difficultés organisationnelles',
    322: 'La gestion simultanée des trois composantes du projet (système embarqué, application web, PCB) a parfois été un défi. Passer de la programmation C++ embarquée au développement JavaScript web, puis à la conception de circuits imprimés, exigeait de la souplesse mais aussi une grande discipline. Pour éviter la dispersion, j\'ai appris à établir un planning précis, à hiérarchiser les tâches et à m\'imposer des délais réalistes.',

    324: 'Au-delà des aspects techniques et organisationnels, ce stage a été très instructif sur le plan humain et social. Il m\'a permis de comprendre comment fonctionne une structure professionnelle et comment les relations humaines influencent directement la réussite d\'un projet. Travailler au sein de l\'ENSPD, aux côtés d\'enseignants-chercheurs, de techniciens et d\'autres stagiaires, m\'a appris l\'importance de la communication, de la patience et du respect mutuel dans un environnement académique et technique. Le projet MERA, par sa dimension multidisciplinaire (médecine, électronique, informatique), m\'a également sensibilisé aux enjeux de la coopération interdisciplinaire et à la nécessité d\'adapter son discours technique en fonction de son interlocuteur.',
}

CONCLUSION = {
    340: 'Ce stage au sein de l\'École Nationale Supérieure Polytechnique de Douala (ENSPD) a été pour moi une expérience riche et formatrice, aussi bien sur le plan académique que professionnel. Il m\'a permis de mettre en pratique les connaissances acquises en formation, tout en découvrant de nouvelles technologies et en développant des compétences que je ne maîtrisais pas auparavant, notamment en développement de systèmes embarqués, en conception de circuits imprimés, en développement web full-stack et en intégration d\'intelligence artificielle.',
    341: 'Le projet MERA, par sa dimension multidisciplinaire, m\'a offert l\'opportunité de travailler sur un projet à fort impact social, visant à améliorer l\'accès aux soins de santé dans les centres médicaux. Au-delà des compétences techniques, ce stage m\'a également appris la rigueur, l\'esprit d\'analyse, le travail en équipe et l\'adaptabilité face aux difficultés rencontrées. Ces qualités sont essentielles dans le domaine de l\'ingénierie et me seront précieuses dans la suite de mon parcours.',
    342: 'En somme, ce stage m\'a offert une vision concrète du métier d\'ingénieur, en me plaçant au cœur d\'un projet innovant et stimulant alliant électronique, intelligence artificielle et santé digitale. Il constitue une étape déterminante dans mon parcours, et renforce ma motivation à poursuivre dans cette voie, afin de contribuer activement au développement technologique et industriel du Cameroun et de la sous-région.',
}

# ============================================================
# APPLICATION
# ============================================================
print('Copie du document original...')
shutil.copy2(SRC, DST)

doc = docx.Document(DST)
print('Paragraphes avant modification:', len(doc.paragraphs))

ALL = {}
for d in (TDM, FIGURES_LIST, INTRO, CH2, CH3, CONCLUSION):
    ALL.update(d)

missing = []
for idx, text in ALL.items():
    if idx < len(doc.paragraphs):
        ok = replace_text_only(doc.paragraphs[idx], text)
        if not ok:
            missing.append(idx)
    else:
        missing.append(idx)

print('Index sans run texte:', missing)


def remove_drawings(para):
    """Retire les images d'un paragraphe (logos ne correspondant pas au projet)."""
    removed = 0
    for r in list(para.runs):
        if 'w:drawing' in r._element.xml or 'w:pict' in r._element.xml:
            r._element.getparent().remove(r._element)
            removed += 1
    return removed


print('Logos Expo/React Native retirés (para 265):', remove_drawings(doc.paragraphs[265]))
print('Légende "Logo Expo" retirée (para 266):', remove_drawings(doc.paragraphs[266]))


def delete_paragraph(para):
    """Supprime complètement un paragraphe du document."""
    el = para._element
    el.getparent().remove(el)


def find_paragraph(predicate):
    for p in doc.paragraphs:
        if predicate(p):
            return p
    return None


# ------------------------------------------------------------
# Suppression de la Figure 12 (interface de diagnostic) :
# légende (paragraphe 300) + image (paragraphe 299) + entrée liste (142).
# Suppression en ordre décroissant pour ne pas décaler les indices.
# ------------------------------------------------------------
for idx in (300, 299, 142):
    delete_paragraph(doc.paragraphs[idx])
print('Figure 12 supprimée (image + légende + entrée liste).')

# ------------------------------------------------------------
# Insertion de la sous-section « Balance numérique intelligente (BLE) »
# (avant le paragraphe « Outils logiciels »)
# ------------------------------------------------------------
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BALANCE_IMG = '/tmp/opencode/new_images/balance_ble.jpeg'
if os.path.exists(BALANCE_IMG):
    ref = find_paragraph(lambda p: p.text.strip() == 'Outils logiciels')
    head = ref.insert_paragraph_before('Balance numérique intelligente (BLE)', style='Body Text')
    desc = ref.insert_paragraph_before(
        'Balance numérique connectée (modèle Famkit) transmettant le poids par Bluetooth Low Energy '
        '(BLE). Le firmware de l\'ESP32 écoute les annonces diffusées par la balance et récupère le '
        'poids en kilogrammes sans aucune liaison filaire. Une mesure est validée lorsque le poids '
        'est compris entre 10 et 300 kg et reste stable à \u00b10,3 kg sur cinq lectures consécutives. '
        'Le poids rejoint ensuite le dossier du patient avec les autres constantes vitales (rythme '
        'cardiaque, SpO2 et température).',
        style='Normal (Web)')
    img = ref.insert_paragraph_before('', style='Normal')
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.add_run().add_picture(BALANCE_IMG, height=Inches(1.4))
    cap = ref.insert_paragraph_before('Figure 5: Balance numérique intelligente (Famkit)', style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    print('Sous-section Balance BLE insérée (3 paragraphes + image).')
else:
    print('ATTENTION: image de la balance introuvable:', BALANCE_IMG)

# ------------------------------------------------------------
# Insertion de l'entrée « Figure 6: Logo Arduino IDE » dans la liste des figures
# (avant l'entrée « Figure 7: Logo TypeScript »)
# ------------------------------------------------------------
ref_list = find_paragraph(lambda p: p.text.startswith('Figure 7: Logo TypeScript'))
new_entry = ref_list.insert_paragraph_before('Figure 6: Logo Arduino IDE\t11', style='table of figures')
print('Entrée liste des figures insérée:', repr(new_entry.text))

doc.save(DST)
doc2 = docx.Document(DST)
print('Paragraphes après modification:', len(doc2.paragraphs))
print('Document enregistré:', DST)
